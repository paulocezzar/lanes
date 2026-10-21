"""Add many people at once from a spreadsheet (CSV or Excel).

Columns are matched by heading, in any order, case-insensitive:
  First name | Last name | Email | Department | Role (Member/Admin)
Only Email and Department are required (plus First name for new accounts).
"""
import csv
import io
import re

from django.contrib.auth import get_user_model

from .models import Workspace, WorkspaceMember

User = get_user_model()
MAX_ROWS = 1000
EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
ALIASES = {
    "first_name": {"first name", "firstname", "first", "forename", "given name"},
    "last_name": {"last name", "lastname", "last", "surname", "family name"},
    "name": {"name", "full name", "fullname", "employee", "employee name"},
    "email": {"email", "e-mail", "email address", "work email"},
    "department": {"department", "dept", "team", "workspace", "site"},
    "role": {"role", "access", "permission"},
}


def _key(heading):
    h = re.sub(r"\s+", " ", str(heading or "").strip().lower().replace("_", " "))
    for key, names in ALIASES.items():
        if h in names:
            return key
    return None


def read_rows(upload):
    """Return a list of dicts with our keys. Raises ValueError with a friendly message."""
    name = (upload.name or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            ws = load_workbook(upload, read_only=True, data_only=True).active
        except Exception:
            raise ValueError("That Excel file couldn't be opened. Try saving it again, or as CSV.")
        table = [["" if c is None else str(c).strip() for c in row] for row in ws.iter_rows(values_only=True)]
    elif name.endswith((".csv", ".txt")):
        raw = upload.read()
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t") if text.strip() else csv.excel
        table = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    else:
        raise ValueError("Upload a CSV or Excel (.xlsx) file.")
    table = [r for r in table if any(c for c in r)]
    if len(table) < 2:
        raise ValueError("The file needs a heading row and at least one person.")
    keys = [_key(h) for h in table[0]]
    if "email" not in keys:
        raise ValueError("Couldn't find an Email column. The first row should have headings like "
                         "First name, Last name, Email, Department, Role.")
    if "department" not in keys:
        raise ValueError("Couldn't find a Department column.")
    rows = []
    for line, r in enumerate(table[1:MAX_ROWS + 1], start=2):
        d = {k: (r[i] if i < len(r) else "") for i, k in enumerate(keys) if k}
        if d.get("name") and not d.get("first_name"):
            parts = d["name"].split()
            d["first_name"], d["last_name"] = parts[0], " ".join(parts[1:])
        d["line"] = line
        rows.append(d)
    if len(table) - 1 > MAX_ROWS:
        raise ValueError(f"That's more than {MAX_ROWS} people. Split the file and upload it in parts.")
    return rows


def plan(rows, create_departments):
    """Work out what each row would do, without changing anything."""
    depts = {w.name.lower(): w for w in Workspace.objects.all()}
    users = {u.email.lower(): u for u in User.objects.filter(email__in=[(r.get("email") or "").strip().lower()
                                                                        for r in rows])}
    seen, out = set(), []
    for r in rows:
        email = (r.get("email") or "").strip().lower()
        dept = re.sub(r"\s+", " ", (r.get("department") or "").strip())[:80]
        role_raw = (r.get("role") or "").strip().lower()
        role = WorkspaceMember.ADMIN if role_raw in ("admin", "administrator", "manager", "department admin") \
            else WorkspaceMember.MEMBER
        item = {"line": r["line"], "email": email, "first": (r.get("first_name") or "").strip()[:150],
                "last": (r.get("last_name") or "").strip()[:150], "dept": dept, "role": role,
                "action": "", "problem": ""}
        user = users.get(email)
        ws = depts.get(dept.lower())
        if not EMAIL.match(email):
            item["problem"] = "Email missing or not valid"
        elif not dept:
            item["problem"] = "No department"
        elif (email, dept.lower()) in seen:
            item["problem"] = "Same person and department listed twice"
        elif not ws and not create_departments:
            item["problem"] = f"No department called “{dept}” (tick “create missing departments”)"
        elif user is None and not item["first"]:
            item["problem"] = "New account needs a first name"
        elif user is not None and not user.is_active:
            item["problem"] = "Account is switched off — switch it back on first"
        elif user is not None and ws and WorkspaceMember.objects.filter(workspace=ws, user=user).exists():
            item["action"] = "already"
        else:
            item["action"] = "new" if user is None else "add"
            item["new_dept"] = ws is None
        seen.add((email, dept.lower()))
        out.append(item)
    return out


def apply(items, actor):
    """Create accounts and memberships. Returns (created_accounts, added_count, new_departments)."""
    from django.db import transaction

    from .views_people import temporary_password
    created, added, new_depts = [], 0, []
    with transaction.atomic():
        depts = {w.name.lower(): w for w in Workspace.objects.all()}
        for it in items:
            if it["action"] not in ("new", "add"):
                continue
            ws = depts.get(it["dept"].lower())
            if ws is None:
                ws = Workspace.objects.create(name=it["dept"])
                depts[ws.name.lower()] = ws
                new_depts.append(ws.name)
            user = User.objects.filter(email__iexact=it["email"]).first()
            if user is None:
                password = temporary_password()
                user = User.objects.create_user(email=it["email"], password=password,
                                                first_name=it["first"], last_name=it["last"])
                created.append({"name": user.display_name, "email": user.email, "password": password,
                                "dept": ws.name})
            elif user.is_guest:
                user.is_guest = False
                user.save(update_fields=["is_guest"])
            _, made = WorkspaceMember.objects.get_or_create(workspace=ws, user=user, defaults={"role": it["role"]})
            added += made
    return created, added, new_depts
